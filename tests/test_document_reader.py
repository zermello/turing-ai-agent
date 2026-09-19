import pymupdf
from docx import Document

import tools.document_reader as dr


def test_txt_reader():
    content = dr.txt_reader("documents/notes.txt")
    assert content == "This is my robotics study plan..."


def test_txt_reader_empty():
    content = dr.txt_reader("documents/empty.txt")
    assert content == ""


def test_txt_reader_missing():
    content = dr.txt_reader("documents/missing.txt")
    assert content == "Path does not exist!"


def test_md_reader():
    content = dr.md_reader("documents/notes.md")
    assert content == '# Robotics\n\n## Languages\n- Python\n- C++\n\n## Goals\nBuild autonomous robots.'


def test_md_reader_missing():
    content = dr.md_reader("documents/missing.md")
    assert content == "Path does not exist!"


def test_md_reader_empty(tmp_path):
    file_path = tmp_path / "empty.md"
    file_path.write_text("", encoding="utf-8")

    assert dr.md_reader(file_path) == ""


def test_json_reader():
    content = dr.json_reader("documents/notes.json")
    assert content == {'name': 'TURING', 'tools': ['calculator', 'weather', 'memory'], 'version': 1}


def test_json_reader_missing():
    content = dr.json_reader("documents/missing.json")
    assert content == "Path doesn't exist"


def test_json_reader_invalid(tmp_path):
    file_path = tmp_path / "invalid.json"
    file_path.write_text('{"name": }', encoding="utf-8")

    assert dr.json_reader(file_path) == "error occured"


def test_csv_reader():
    content = dr.csv_reader("documents/notes.csv")

    assert content == [{'name': 'TURING',
                        'type': 'AI Agent',
                        'language': 'Python',
                        'status': 'Active'},
                        {'name': 'RoboArm',
                        'type': 'Manipulator',
                        'language': 'C++',
                        'status': 'Development'},
                        {'name': 'VisionBot',
                        'type': 'Computer Vision',
                        'language': 'Python',
                        'status': 'Testing'},
                        {'name': 'Atlas',
                        'type': 'Humanoid',
                        'language': 'C++',
                        'status': 'Research'},
                        {'name': 'DroneX',
                        'type': 'Autonomous Drone',
                        'language': 'Python',
                        'status': 'Prototype'}]


def test_csv_reader_missing():
    content = dr.csv_reader("documents/missing.csv")
    assert content == "PATH DOES NOT EXIST!"


def test_csv_reader_empty(tmp_path):
    file_path = tmp_path / "empty.csv"
    file_path.write_text("", encoding="utf-8")

    assert dr.csv_reader(file_path) == []


def test_pdf_reader():
    content = dr.pdf_reader("documents/notes.pdf")

    assert isinstance(content, list)
    assert any("robotics" in page.lower() for page in content)


def test_pdf_reader_missing():
    content = dr.pdf_reader("documents/missing.pdf")
    assert content == "PATH DOESN'T EXIST!!"


def test_pdf_reader_empty(tmp_path):
    file_path = tmp_path / "empty.pdf"
    file_path.write_bytes(b"")

    assert dr.pdf_reader(file_path) == "Error occured"


def test_docx_reader():
    content = dr.docx_reader("documents/notes.docx")

    assert isinstance(content, list)
    assert any("robotics" in paragraph.lower() for paragraph in content)


def test_docx_reader_missing():
    content = dr.docx_reader("documents/missing.docx")
    assert content == "PATH DOESN'T EXIST!!"


def test_docx_reader_empty(tmp_path):
    file_path = tmp_path / "empty.docx"
    Document().save(file_path)

    assert dr.docx_reader(file_path) == []


def test_read_document_unsupported_extension(tmp_path):
    file_path = tmp_path / "notes.rtf"
    file_path.write_text("content", encoding="utf-8")

    assert dr.read_document(file_path) == "Unsupported file type"


def test_read_document_missing_file(tmp_path):
    file_path = tmp_path / "missing.txt"

    assert dr.read_document(file_path) == "Path does not exist!"


def test_readers_support_unicode(tmp_path):
    txt_path = tmp_path / "unicode.txt"
    md_path = tmp_path / "unicode.md"
    json_path = tmp_path / "unicode.json"
    csv_path = tmp_path / "unicode.csv"
    pdf_path = tmp_path / "unicode.pdf"
    docx_path = tmp_path / "unicode.docx"

    txt_path.write_text("こんにちは, café, робот", encoding="utf-8")
    md_path.write_text("# Привет 世界", encoding="utf-8")
    json_path.write_text('{"message": "Olá 世界"}', encoding="utf-8")
    csv_path.write_text("message\nこんにちは 世界\n", encoding="utf-8")

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "café")
    pdf.save(pdf_path)
    pdf.close()

    document = Document()
    document.add_paragraph("Olá 世界")
    document.save(docx_path)

    assert dr.txt_reader(txt_path) == "こんにちは, café, робот"
    assert dr.md_reader(md_path) == "# Привет 世界"
    assert dr.json_reader(json_path) == {"message": "Olá 世界"}
    assert dr.csv_reader(csv_path) == [{"message": "こんにちは 世界"}]
    assert any("café" in page.lower() for page in dr.pdf_reader(pdf_path))
    assert dr.docx_reader(docx_path) == ["Olá 世界"]